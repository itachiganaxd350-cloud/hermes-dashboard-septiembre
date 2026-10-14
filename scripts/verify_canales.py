# -*- coding: utf-8 -*-
"""Verificacion de la seccion CANALES: dataset.json vs canales 2026.xlsx (celdas)"""
import openpyxl, json

CAN = r"C:\Users\ASUS\Desktop\canales 2026.xlsx"
DS = r"C:\Users\ASUS\Documents\cierre-septiembre-2026\src\data\dataset.json"
ds = json.load(open(DS, encoding="utf-8"))["canales"]
wb = openpyxl.load_workbook(CAN, data_only=True)
ws = wb["INFORM CANALES "]
COLS = {2022: 5, 2023: 8, 2024: 11, 2025: 12, 2026: 13}
fails, checks = [], 0


def cmp(name, ex, dv):
    global checks
    checks += 1
    ok = (ex is None and dv is None) or (isinstance(ex, (int, float)) and isinstance(dv, (int, float)) and abs(ex - dv) < 1e-9)
    if not ok:
        fails.append(f"{name}: excel={ex!r} ds={dv!r}")


for key, r0 in [("terrazas", 3), ("comercializadora", 21)]:
    for y, col in COLS.items():
        for i in range(12):
            cmp(f"{key}[{y}] fila{i+1}", ws.cell(row=r0 + i, column=col).value, ds[key][str(y)][i])
# totales
for y, col in COLS.items():
    cmp(f"total terrazas {y}", ws.cell(row=15, column=col).value, ds["totales"]["terrazas"][str(y)])
    cmp(f"total comerc {y}", ws.cell(row=33, column=col).value, ds["totales"]["comercializadora"][str(y)])
# participacion
for i in range(12):
    cmp(f"participacion fila{i+1}", ws.cell(row=21 + i, column=14).value, ds["participacion2026"][i])
cmp("participacion total", ws.cell(row=33, column=14).value, ds["totales"]["participacion2026"])
# sumas internas 2026 (Ene-Sep)
t26 = sum(v for v in ds["terrazas"]["2026"] if v)
c26 = sum(v for v in ds["comercializadora"]["2026"] if v)
cmp("suma terrazas 2026", ds["totales"]["terrazas"]["2026"], t26)
cmp("suma comerc 2026", ds["totales"]["comercializadora"]["2026"], c26)
print(f"CANALES CHECKS: {checks} | FALLOS: {len(fails)}")
print("RESULTADO:", "100% FIEL" if not fails else "REVISAR")
for f in fails:
    print(" -", f)
print(f"Terrazas 2026 Ene-Sep: {t26} | Comercializadora 2026 Ene-Sep: {c26}")
print(f"Participacion total: {ds['totales']['participacion2026']:.4f}% | septiembre: {ds['participacion2026'][8]:.4f}%")
